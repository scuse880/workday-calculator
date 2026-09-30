from collections import Counter, defaultdict
from pathlib import Path
import json
import re
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent / "_vendor"))
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "testsource"

def cell_kind(v):
    if v is None:
        return "blank"
    return type(v).__name__

def workbook_summary(path):
    wb = load_workbook(path, read_only=False, data_only=True, keep_vba=path.suffix.lower()==".xlsm")
    items = []
    for ws in wb.worksheets:
        nonempty = 0
        for row in ws.iter_rows():
            nonempty += sum(c.value is not None for c in row)
        items.append({"title":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,
                      "nonempty":nonempty,"merged_ranges":len(ws.merged_cells.ranges),
                      "hidden":ws.sheet_state})
    return {"file":path.name,"bytes":path.stat().st_size,"sheets":items}

allbooks=[]
for path in sorted(SRC.glob("*.xlsx")):
    if not path.name.startswith("~$"):
        allbooks.append(workbook_summary(path))
allbooks.append(workbook_summary(ROOT / "workbook" / "근무일수계산.xlsm"))

out={"books":allbooks}
over=[]
for label in ("초과근무 월별집계(생년월일 표시).xlsx", "초과근무 월별집계(생년월일 미표시).xlsx"):
    ws=load_workbook(SRC/label,read_only=False,data_only=True).worksheets[0]
    r=5; cells=[]; hops=[]
    while r<=ws.max_row:
        value=ws.cell(r,3).value
        if value is None or not str(value).strip():
            hops.append(r)
            r+=3
            value=ws.cell(r,3).value
            if value is None or not str(value).strip():
                break
        cells.append((r,str(value).strip()))
        r+=1
    names=defaultdict(list); bad=[]
    for row,value in cells:
        match=re.fullmatch(r"(.+?)\s*\((\d{6})\)",value)
        if match:
            names[match.group(1).strip()].append(match.group(2))
        else:
            if "(" in value or ")" in value:bad.append(row)
            names[value].append("")
    dups={name:dates for name,dates in names.items() if len(dates)>1}
    over.append({"file":label,"employee_cells":len(cells),"unique_names":len(names),
                 "birthdate_count":sum(bool(x) for dates in names.values() for x in dates),
                 "duplicate_name_groups":len(dups),"duplicate_name_people":sum(len(d) for d in dups.values()),
                 "duplicate_name_missing_birth_count":sum(not x for d in dups.values() for x in d),
                 "duplicate_name_identical_birth_group_count":sum(len(d)!=len(set(d)) for d in dups.values()),
                 "invalid_parentheses_rows":bad,"gap_hops":len(hops),"termination_rows": [r-3,r],
                 "row_range":[cells[0][0],cells[-1][0]] if cells else None})
out["overtime"]=over
birth_ws=load_workbook(SRC/"초과근무 월별집계(생년월일 표시).xlsx",read_only=False,data_only=True).worksheets[0]
employee_names=set()
duplicate_names=set()
name_counts=Counter()
for row in range(5,birth_ws.max_row+1):
    value=birth_ws.cell(row,3).value
    if value is None or not str(value).strip():continue
    name=re.sub(r"\s*\(\d{6}\)$","",str(value).strip()).strip()
    name_counts[name]+=1
    employee_names.add(name)
duplicate_names={name for name,n in name_counts.items() if n>1}

ws=load_workbook(SRC/"근무상황목록.xlsx",read_only=False,data_only=True).worksheets[0]
status=[];r=2
while r<=ws.max_row and ws.cell(r,3).value is not None and str(ws.cell(r,3).value).strip():
    c=ws.cell(r,3).value;d=ws.cell(r,4).value;e=ws.cell(r,5).value;g=ws.cell(r,7).value
    status.append((r,str(c),str(d or ""),str(e or ""),str(g or "")))
    r+=1
out["work_status"]={"rows":len(status),"stop_row":r,
                    "newline_id_rows":sum("\n" in c or "\r" in c for _,c,_,_,_ in status),
                    "skip_reason_rows":sum("육아시간" in g or "모성보호시간" in g for _,_,_,_,g in status),
                    "period_tilde_counts":dict(Counter(e.count("~") for _,_,_,e,_ in status)),
                    "cell_value_kinds":dict(Counter(cell_kind(ws.cell(row,5).value) for row,_,_,_,_ in status)),
                    "first_last_rows":[status[0][0],status[-1][0]] if status else None}
status_eligible=[(row,c,d,e,g) for row,c,d,e,g in status if "육아시간" not in g and "모성보호시간" not in g]
status_target=[(row,c,d,e,g) for row,c,d,e,g in status_eligible if c.splitlines()[0].strip() in employee_names]
status_duplicated=[(row,c,d,e,g) for row,c,d,e,g in status_target if c.splitlines()[0].strip() in duplicate_names]
out["work_status"]["eligible_rows"]=len(status_eligible)
out["work_status"]["matching_employee_rows"]=len(status_target)
out["work_status"]["nonmatching_employee_rows"]=len(status_eligible)-len(status_target)
out["work_status"]["duplicate_name_rows"]=len(status_duplicated)
out["work_status"]["duplicate_name_distinct_id_count"]=len(set(c.splitlines()[1].strip().strip("()") for _,c,_,_,_ in status_duplicated if len(c.splitlines())>1))
out["work_status"]["duplicate_name_id_shape_counts"]=dict(Counter(bool(re.fullmatch(r"[Cc]\d{9}",c.splitlines()[1].strip().strip("()"))) if len(c.splitlines())>1 else False for _,c,_,_,_ in status_duplicated))
out["work_status"]["period_year_month_counts"]=dict(Counter(re.search(r"\d{4}-\d{1,2}",e).group(0) if re.search(r"\d{4}-\d{1,2}",e) else "unknown" for _,_,_,e,_ in status_target))

ws=load_workbook(SRC/"출장 근무상황부 개인별.xlsx",read_only=False,data_only=True).worksheets[0]
trip=[];r=5;consecutive=0;skipped=[]
while r<=ws.max_row:
    b=ws.cell(r,2).value;c=ws.cell(r,3).value;d=ws.cell(r,4).value
    if all(v is not None and str(v).strip() for v in (b,c,d)) and "직급" not in str(c):
        consecutive=0
        trip.append((r,str(b).strip(),c,d))
    else:
        consecutive+=1;skipped.append(r)
        if consecutive==5:break
    r+=1
out["trip"]={"target_rows":len(trip),"first_last_rows":[trip[0][0],trip[-1][0]] if trip else None,
             "stop_row":r,"skipped_rows":skipped[:30],"masked_name_rows":sum("(" in b for _,b,_,_ in trip),
             "start_kinds":dict(Counter(cell_kind(c) for _,_,c,_ in trip)),
             "end_kinds":dict(Counter(cell_kind(d) for _,_,_,d in trip)),
             "malformed_masked_rows":[row for row,b,_,_ in trip if "(" in b and not re.fullmatch(r".+\d{2}\s*\([^*\s()]{3}\*+\)",b)]}
all_trip_targets=[]
trip_row_shape=[]
for row in range(5,ws.max_row+1):
    b=ws.cell(row,2).value;c=ws.cell(row,3).value;d=ws.cell(row,4).value
    is_target=all(v is not None and str(v).strip() for v in (b,c,d)) and "직급" not in str(c)
    if is_target:all_trip_targets.append(row)
    if row <= 45:
        trip_row_shape.append({"row":row,"b":bool(b),"c":bool(c),"d":bool(d),
                               "c_type":cell_kind(c),"c_has_grade":bool(c and "직급" in str(c)),
                               "target":is_target})
out["trip"]["all_target_rows_count"]=len(all_trip_targets)
out["trip"]["later_target_rows_count"]=sum(row>r for row in all_trip_targets)
out["trip"]["first_later_target_rows"]= [row for row in all_trip_targets if row>r][:20]
out["trip"]["last_target_row"]=all_trip_targets[-1]
out["trip"]["max_non_target_gap_between_targets"]=max((b-a-1 for a,b in zip(all_trip_targets,all_trip_targets[1:])),default=0)
out["trip"]["non_target_gap_distribution"]=dict(Counter(b-a-1 for a,b in zip(all_trip_targets,all_trip_targets[1:])))
out["trip"]["masked_all_target_count"]=sum("(" in str(ws.cell(row,2).value) for row in all_trip_targets)
out["trip"]["all_target_start_kinds"]=dict(Counter(cell_kind(ws.cell(row,3).value) for row in all_trip_targets))
out["trip"]["all_target_end_kinds"]=dict(Counter(cell_kind(ws.cell(row,4).value) for row in all_trip_targets))
out["trip"]["remaining_nonempty_after_last_target"]=sum(any(ws.cell(row,col).value is not None for col in (2,3,4)) for row in range(all_trip_targets[-1]+1,ws.max_row+1))
trip_names=[]
for row in all_trip_targets:
    raw=str(ws.cell(row,2).value).strip()
    match=re.fullmatch(r"(.+?)\d{2}\s*\([^*\s()]{3}\*+\)",raw)
    trip_names.append((row,match.group(1).strip() if match else raw, bool(match)))
out["trip"]["matching_employee_rows"]=sum(name in employee_names for _,name,_ in trip_names)
out["trip"]["nonmatching_employee_rows"]=sum(name not in employee_names for _,name,_ in trip_names)
out["trip"]["modal_rows_by_known_employee"]=sum(name in employee_names and (masked or name in duplicate_names) for _,name,masked in trip_names)
out["trip"]["record_start_year_month_counts"]=dict(Counter(re.search(r"\d{4}\.\d{1,2}",str(ws.cell(row,3).value)).group(0) if re.search(r"\d{4}\.\d{1,2}",str(ws.cell(row,3).value)) else "unknown" for row in all_trip_targets))
out["trip"]["first_45_row_shapes"]=trip_row_shape

wb=load_workbook(ROOT/"workbook"/"근무일수계산.xlsm",read_only=False,data_only=True,keep_vba=True)
base=wb["작업"]
vac=wb["방학중근무"]
vacrows=[];r=3
while r<=vac.max_row:
    a=vac.cell(r,1).value
    if a is None or not str(a).strip():break
    vacrows.append((r,a,vac.cell(r,2).value,vac.cell(r,3).value,vac.cell(r,4).value));r+=1
out["actual_xlsm"]={"base_settings":{"B3":str(base["B3"].value),"D3":str(base["D3"].value),
                                    "B5":str(base["B5"].value),"D5":str(base["D5"].value),
                                    "B6":str(base["B6"].value),"D6":str(base["D6"].value)},
                    "vacation_rows":len(vacrows),"vacation_stop_row":r,
                    "vacation_value_kinds":{"date":dict(Counter(cell_kind(x[2]) for x in vacrows)),
                                            "start":dict(Counter(cell_kind(x[3]) for x in vacrows)),
                                            "end":dict(Counter(cell_kind(x[4]) for x in vacrows))},
                    "vacation_birthdate_rows":sum("(" in str(x[1]) for x in vacrows),
                    "vacation_first_last_rows":[vacrows[0][0],vacrows[-1][0]] if vacrows else None}
out["actual_xlsm"]["vacation_matching_employee_rows"]=sum(str(x[1]).strip() in employee_names for x in vacrows)
out["actual_xlsm"]["vacation_nonmatching_employee_rows"]=sum(str(x[1]).strip() not in employee_names for x in vacrows)
out["actual_xlsm"]["vacation_date_year_month_counts"]=dict(Counter(x[2].strftime("%Y-%m") if hasattr(x[2],"strftime") else "invalid" for x in vacrows))

print(json.dumps(out,ensure_ascii=True,indent=2))
