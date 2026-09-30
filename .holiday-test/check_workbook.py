from openpyxl import load_workbook
from pathlib import Path

p = Path(__file__).resolve().parents[1] / 'workbook' / '근무일수계산.xlsm'
w = load_workbook(p, read_only=True, keep_vba=True, data_only=False)
print('sheets:', w.sheetnames)
print('work month:', w['작업']['B3'].value, w['작업']['D3'].value)
print('holidays:', list(w['공휴일'].values)[:8])
